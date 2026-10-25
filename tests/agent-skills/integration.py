#!/usr/bin/env python3
"""Networked native smoke test in a new HOME; no host configuration or activation."""
import argparse
import base64
import hashlib
import json
import os
from pathlib import Path
import subprocess
import time
import zipfile

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--package", type=Path, required=True)
parser.add_argument("--uv", type=Path, required=True)
parser.add_argument("--node", type=Path, required=True)
parser.add_argument("--pdftoppm", type=Path, required=True)
parser.add_argument("--output", type=Path, required=True, help="New disposable directory; must not already exist")
args = parser.parse_args()
out = args.output.resolve(); out.mkdir(parents=True, exist_ok=False)
home = out / "home"; home.mkdir()
tmp = out / "tmp"; tmp.mkdir()
# No inherited secrets, shell configuration, npmrc, uv config or browser state.
env = {"HOME": str(home), "PATH": "/usr/bin:/bin", "TMPDIR": str(tmp),
       "XDG_CACHE_HOME": str(home / ".cache"), "UV_CACHE_DIR": str(home / ".cache/uv"),
       "UV_PYTHON_INSTALL_DIR": str(home / ".local/share/uv/python"),
       "UV_TOOL_DIR": str(home / ".local/share/uv/tools"), "UV_TOOL_BIN_DIR": str(home / ".local/bin"),
       "PLAYWRIGHT_BROWSERS_PATH": str(home / ".cache/ms-playwright"),
       "npm_config_cache": str(home / ".cache/npm"), "CI": "1"}
setup = args.package / "bin/agent-skills-setup"
python = args.package / "bin/agent-skills-python"
root = home / ".local/share/agent-skills"
venv = root / "venv"
commands = []


def run(label, command, extra_env=None):
    command = [str(item) for item in command]
    start = time.monotonic()
    with (out / (label + ".log")).open("w") as log:
        result = subprocess.run(command, env=env | (extra_env or {}), cwd=out,
                                stdout=log, stderr=subprocess.STDOUT, timeout=900)
    commands.append({"label": label, "command": command, "exit_code": result.returncode,
                     "seconds": time.monotonic() - start})
    (out / "commands.json").write_text(json.dumps(commands, indent=2) + "\n")
    if result.returncode:
        raise RuntimeError(f"{label} failed; see {out / (label + '.log')}")


run("initial-setup", [setup])
inode = venv.stat().st_ino
config = (venv / "pyvenv.cfg").read_bytes()
# A local wheel makes extras recovery independent of an unrelated public package.
wheels = out / "wheels"; wheels.mkdir()
wheel = wheels / "skills_retained_fixture-1.0-py3-none-any.whl"
files = {"skills_retained_fixture.py": b"VALUE = 'retained'\n",
         "skills_retained_fixture-1.0.dist-info/METADATA":
         b"Metadata-Version: 2.1\nName: skills-retained-fixture\nVersion: 1.0\n",
         "skills_retained_fixture-1.0.dist-info/WHEEL":
         b"Wheel-Version: 1.0\nGenerator: fixture\nRoot-Is-Purelib: true\nTag: py3-none-any\n"}
record = ""
for name, data in files.items():
    digest = base64.urlsafe_b64encode(hashlib.sha256(data).digest()).decode().rstrip("=")
    record += f"{name},sha256={digest},{len(data)}\n"
files["skills_retained_fixture-1.0.dist-info/RECORD"] = (
    record + "skills_retained_fixture-1.0.dist-info/RECORD,,\n"
).encode()
with zipfile.ZipFile(wheel, "w") as archive:
    for name, data in files.items(): archive.writestr(name, data)
env["UV_FIND_LINKS"] = str(wheels)
run("add-python-extra", [args.uv, "pip", "install", "--python", venv / "bin/python", "--no-index", wheel])
extra_node = root / "local-test-package"; extra_node.mkdir()
(extra_node / "package.json").write_text(json.dumps({
    "name": "skills-retained-fixture", "version": "1.0.0", "main": "index.js"
}))
(extra_node / "index.js").write_text("module.exports = 'retained';\n")
npm = args.node.with_name("npm")
node_env = {"PATH": str(args.node.parent) + ":/usr/bin:/bin"}
run("add-node-extra", [npm, "install", "--prefix", root / "node", "--ignore-scripts",
                       "--no-audit", "--no-fund", extra_node], node_env)
run("repeat-setup", [setup])
assert venv.stat().st_ino == inode, "valid venv was recreated"
assert (venv / "pyvenv.cfg").read_bytes() == config, "valid interpreter binding changed"
# Force only the binding test to fail; the recoverable installed metadata stays intact.
(venv / "pyvenv.cfg").write_text(config.decode().replace("home = ", "home = /synthetic-wrong-binding/", 1))
run("recover-binding", [setup, "venv"])
assert not (root / "venv.old").exists(), "successful recovery left old venv"
assert venv.stat().st_ino != inode, "wrong-bound venv was not rebuilt"
probe = out / "documents.py"
probe.write_text("""import json,sys
from pathlib import Path
import skills_retained_fixture
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen.canvas import Canvas
from pypdf import PdfReader
assert skills_retained_fixture.VALUE == "retained"
book = Workbook(); book.active["A1"] = "Skill setup"; book.active["B1"] = 42
book.save("fixture.xlsx")
assert load_workbook("fixture.xlsx").active["B1"].value == 42
pdf = Canvas("fixture.pdf"); pdf.drawString(72, 750, "Agent skills setup smoke test"); pdf.save()
assert "Agent skills setup smoke test" in PdfReader("fixture.pdf").pages[0].extract_text()
print(json.dumps({"python":sys.version,"prefix":sys.prefix,"base_prefix":sys.base_prefix}))
""")
run("documents-python", [python, probe])
js = out / "documents.cjs"
js.write_text("""const fs = require('fs');
const { Document, Packer, Paragraph } = require('docx');
if (require('skills-retained-fixture') !== 'retained') throw Error('extra missing');
const doc = new Document({sections:[{children:[new Paragraph('Agent skills document smoke')]}]});
Packer.toBuffer(doc).then(b => fs.writeFileSync('fixture.docx',b)).catch(e => {console.error(e);process.exitCode=1;});
""")
run("documents-node", [args.node, js], {"NODE_PATH": str(root / "node/node_modules")})
with zipfile.ZipFile(out / "fixture.docx") as archive:
    assert b"Agent skills document smoke" in archive.read("word/document.xml")
run("render-pdf", [args.pdftoppm, "-singlefile", "-png", "-r", "72", out / "fixture.pdf", out / "rendered"])
assert (out / "rendered.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
html = out / "fixture.html"
html.write_text('<h1>Agent skills browser smoke</h1><button onclick="this.textContent=\'Passed\'">Run check</button>')
browser = out / "browser.py"
browser.write_text("""from pathlib import Path
from playwright.sync_api import sync_playwright, expect
with sync_playwright() as p:
    browser = p.chromium.launch(headless=True)
    try:
        page = browser.new_page()
        page.goto(Path("fixture.html").resolve().as_uri())
        expect(page.get_by_role("heading", name="Agent skills browser smoke")).to_be_visible()
        page.get_by_role("button", name="Run check").click()
        expect(page.get_by_role("button", name="Passed")).to_be_visible()
        page.screenshot(path="browser.png", full_page=True)
        print(browser.version)
    finally:
        browser.close()
""")
run("browser", [python, browser])
run("python-packages", [args.uv, "pip", "list", "--python", venv / "bin/python", "--format", "json"])
run("node-packages", [npm, "list", "--prefix", root / "node", "--depth=0", "--json"], node_env)
run("markitdown", [home / ".local/bin/markitdown", out / "fixture.xlsx"])
run("uv-version", [args.uv, "--version"])
run("node-version", [args.node, "--version"])
(out / "result.json").write_text(json.dumps({
    "passed": True, "scope": "native disposable HOME; no VM or host activation",
    "valid_venv_preserved": True, "recovery_extras_preserved": True,
    "pdf_render": "rendered.png", "browser_screenshot": "browser.png",
    "package": str(args.package), "home": str(home),
}, indent=2) + "\n")
print(out)
